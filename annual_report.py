"""Génération du rapport annuel d'activité de la structure utilisatrice."""

from __future__ import annotations

import io
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle


BLUE = "#2e718f"
BLUE_DARK = "#1f5065"
RED = "#e30613"
LIGHT = "#eef3f4"
TEXT = "#20282b"


def _paragraphs(value):
    return [line.strip() for line in (value or "").splitlines() if line.strip()]


def _metric_rows(statistics):
    summary = statistics["selected_summary"]
    services = statistics["service_statistics"]
    return [
        ("Passages usagers", summary["sessions"]),
        ("Visiteurs", summary["visitors"]),
        ("Usagers distincts", summary["unique_users"]),
        ("Temps de présence", f"{str(summary['hours']).replace('.', ',')} h"),
        ("Animations", services["animations"]),
        ("Participants aux animations", services["animation_participants"]),
        ("Durée des animations", _minutes_label(services["animation_duration_minutes"])),
        ("Créneaux réservables", services["reservations"]),
        ("Participants aux réservations", services["reservation_participants"]),
        ("Durée des réservations", _minutes_label(services["reservation_duration_minutes"])),
        ("Locations de machines", services["rentals"]),
        ("Recettes réservations et locations", f"{services['revenue']} €"),
    ]


def _minutes_label(value):
    hours, minutes = divmod(max(0, int(value or 0)), 60)
    if hours:
        return f"{hours} h {minutes:02d} min"
    return f"{minutes} min"


def _distribution_text(rows, limit=8):
    selected = [row for row in rows if row.get("count")][:limit]
    return " · ".join(
        f"{row['label']} : {row['count']} ({str(row['percentage']).replace('.', ',')} %)"
        for row in selected
    ) or "Aucune donnée renseignée."


def _structure_identity(structure):
    structure = structure or {"name": "Mon FabLab", "description": "Espace de fabrication numérique", "address": ""}
    name = structure.get("name") or "FabLab"
    description = structure.get("description") or ""
    address = structure.get("address") or ""
    subtitle = " · ".join(part for part in (name, description) if part)
    footer = " - ".join(part for part in (name, description) if part)
    if address:
        footer += " · " + address
    return name, subtitle, footer


def generate_activity_report_docx(report, statistics, logo_path: Path | None = None,
                                  structure=None) -> bytes:
    name, subtitle_text, footer_text = _structure_identity(structure)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT.lstrip("#"))
    normal.paragraph_format.space_after = Pt(2)
    for style_name, size, color in (
        ("Title", 21, BLUE_DARK), ("Heading 1", 12.5, BLUE_DARK), ("Heading 2", 9, RED)
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)
    title.add_run(f"Rapport d’activité {report['year']}")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(subtitle_text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(RED.lstrip("#"))

    if report.get("introduction"):
        document.add_heading("Présentation de l’année", level=1)
        for text in _paragraphs(report["introduction"]):
            document.add_paragraph(text)

    document.add_heading("Chiffres clés", level=1)
    metrics = _metric_rows(statistics)
    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Light Shading Accent 1"
    for index in range(0, len(metrics), 4):
        cells = table.add_row().cells
        for cell in cells:
            cell.width = Cm(4.55)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for cell, metric in zip(cells, metrics[index:index + 4]):
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(metric[1]))
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor.from_string(BLUE.lstrip("#"))
            label = paragraph.add_run(f"\n{metric[0]}")
            label.font.size = Pt(7.5)

    for field, heading in (
        ("highlights", "Temps forts et animations"),
        ("new_equipment", "Nouveaux équipements"),
        ("changes", "Fonctionnement, changements et nouveautés"),
        ("partnerships", "Partenariats et projets"),
    ):
        if report.get(field):
            document.add_heading(heading, level=1)
            for text in _paragraphs(report[field]):
                document.add_paragraph(text)

    document.add_heading("Profil des publics", level=1)
    profile = document.add_table(rows=2, cols=2)
    profile.alignment = WD_TABLE_ALIGNMENT.CENTER
    profile.autofit = False
    for cell, (heading, rows) in zip(
        [cell for row in profile.rows for cell in row.cells],
        (
            ("Genres", statistics["gender_rows"]),
            ("Tranches d’âge", statistics["age_rows"]),
            ("Communes principales", statistics["city_rows"]),
            ("Pays de nationalité", statistics["nationality_rows"]),
        ),
    ):
        cell.width = Cm(9.1)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(1)
        heading_run = paragraph.add_run(heading)
        heading_run.bold = True
        heading_run.font.size = Pt(8.5)
        heading_run.font.color.rgb = RGBColor.from_string(RED.lstrip("#"))
        values = cell.add_paragraph(_distribution_text(rows))
        values.paragraph_format.space_after = Pt(0)
        for run in values.runs:
            run.font.size = Pt(7.5)

    if report.get("additional_notes"):
        document.add_page_break()
        document.add_heading("Compléments", level=1)
        for text in _paragraphs(report["additional_notes"]):
            document.add_paragraph(text)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(footer_text)
    footer_run.font.size = Pt(7)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def generate_activity_report_pdf(report, statistics, logo_path: Path | None = None,
                                 structure=None) -> bytes:
    name, subtitle_text, footer_text = _structure_identity(structure)
    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=10 * mm, bottomMargin=12 * mm,
        title=f"Rapport d'activité {report['year']} {name}",
        author=name,
    )
    sample = getSampleStyleSheet()
    title = ParagraphStyle(
        "ReportTitle", parent=sample["Title"], fontName="Helvetica-Bold",
        fontSize=20, leading=23, textColor=colors.HexColor(BLUE_DARK), alignment=TA_CENTER,
    )
    h1 = ParagraphStyle(
        "ReportH1", parent=sample["Heading1"], fontName="Helvetica-Bold",
        fontSize=12, leading=14, spaceBefore=3.5 * mm, spaceAfter=1.5 * mm,
        textColor=colors.HexColor(BLUE_DARK),
    )
    h2 = ParagraphStyle(
        "ReportH2", parent=sample["Heading2"], fontName="Helvetica-Bold",
        fontSize=8.5, leading=10, textColor=colors.HexColor(RED), spaceBefore=1.5 * mm,
    )
    body = ParagraphStyle(
        "ReportBody", parent=sample["BodyText"], fontName="Helvetica",
        fontSize=8.2, leading=10.5, textColor=colors.HexColor(TEXT), spaceAfter=1.5 * mm,
    )
    story = []
    story.extend([
        Paragraph(f"Rapport d’activité {report['year']}", title),
        Paragraph(escape(subtitle_text), ParagraphStyle("Sub", parent=h2, alignment=TA_CENTER, fontSize=10)),
    ])
    if report.get("introduction"):
        story.append(Paragraph("Présentation de l’année", h1))
        story.extend(Paragraph(text, body) for text in _paragraphs(report["introduction"]))

    story.append(Paragraph("Chiffres clés", h1))
    metrics = _metric_rows(statistics)
    metric_cells = []
    for label, value in metrics:
        metric_cells.append(Paragraph(f"<font size='12.5' color='{BLUE}'><b>{value}</b></font><br/><font size='7.2'>{label}</font>", ParagraphStyle("Metric", parent=body, alignment=TA_CENTER, leading=12.5)))
    metric_data = [metric_cells[index:index + 4] for index in range(0, len(metric_cells), 4)]
    metric_table = Table(metric_data, colWidths=[46.5 * mm] * 4)
    metric_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT)),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9e2e4")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9e2e4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
    ]))
    story.append(metric_table)

    for field, heading in (
        ("highlights", "Temps forts et animations"),
        ("new_equipment", "Nouveaux équipements"),
        ("changes", "Fonctionnement, changements et nouveautés"),
        ("partnerships", "Partenariats et projets"),
    ):
        if report.get(field):
            story.append(Paragraph(heading, h1))
            story.extend(Paragraph(text, body) for text in _paragraphs(report[field]))

    story.append(Paragraph("Profil des publics", h1))
    profile_cells = []
    for heading, rows in (
        ("Genres", statistics["gender_rows"]),
        ("Tranches d’âge", statistics["age_rows"]),
        ("Communes principales", statistics["city_rows"]),
        ("Pays de nationalité", statistics["nationality_rows"]),
    ):
        profile_cells.append(Paragraph(
            f"<font color='{RED}'><b>{heading}</b></font><br/>{_distribution_text(rows)}",
            ParagraphStyle("Profile", parent=body, fontSize=7.5, leading=9),
        ))
    profile_table = Table(
        [profile_cells[:2], profile_cells[2:]], colWidths=[93 * mm, 93 * mm]
    )
    profile_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#d9e2e4")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d9e2e4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(profile_table)
    if report.get("additional_notes"):
        story.append(PageBreak())
        story.append(Paragraph("Compléments", h1))
        story.extend(Paragraph(text, body) for text in _paragraphs(report["additional_notes"]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#69777c"))
        canvas.drawCentredString(A4[0] / 2, 8 * mm, f"{footer_text} · page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
