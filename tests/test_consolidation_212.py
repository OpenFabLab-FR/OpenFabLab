"""Fictional statistical, export and read-only health acceptance scenarios."""
import csv,io,json,os,re,sqlite3,subprocess,sys,unittest
from datetime import datetime,timezone
from unittest import mock
from pathlib import Path
from tests.test_families import FamilyTests
import app
import statistics_quality as quality
import installation_health as health
import animation_attendance as attendance
from visual_report import public_summary,generate_visual_report
from annual_report import _metric_rows,generate_activity_report_pdf,generate_activity_report_docx

def pdf_text(content):
    python=os.environ.get('PDF_PYTHON')
    if not python:raise unittest.SkipTest('PDF_PYTHON with pypdf required for independent text extraction')
    result=subprocess.run([python,'-c','import io,sys;from pypdf import PdfReader;print("\\n".join(p.extract_text() for p in PdfReader(io.BytesIO(sys.stdin.buffer.read())).pages))'],input=content,capture_output=True,check=True)
    return result.stdout.decode()

class ConsolidationTests(unittest.TestCase):
    def setUp(self):
        self.f=FamilyTests();self.f.setUp();self.addCleanup(self.f.doCleanups)
        self.db=self.f.db;self.client=self.f.client;self.service=self.f.service
        self.now=datetime(2026,10,21,12,tzinfo=timezone.utc)

    def past(self,actual=7):
        with self.db:self.db.execute("UPDATE fablab_services SET service_date='2025-08-14',actual_participants=? WHERE id=?",(actual,self.service))

    def stats(self,year=2025):return app.load_statistics(self.db,year)

    def test_future_animations_and_prefilled_participation_excluded(self):
        with self.db:self.db.execute("UPDATE fablab_services SET service_date='2200-01-01',actual_participants=17")
        value=quality.service_statistics(self.db,'total',self.now)
        self.assertEqual((value['animations'],value['animation_participants'],value['animations_scheduled']),(0,0,1))

    def test_past_does_not_claim_held_or_use_capacity(self):
        self.past(7);value=quality.service_statistics(self.db,2025,self.now)
        self.assertEqual((value['animations'],value['animation_participants']),(1,7))
        self.assertEqual(value['animation_presence_incomplete'],0)
        self.assertNotIn('held',value)

    def test_missing_legacy_zero_is_not_a_certain_zero(self):
        self.past(0);self.assertEqual(self.stats()['service_statistics']['animation_presence_incomplete'],1)

    def test_today_before_and_at_end_in_structure_timezone(self):
        with self.db:
            app.write_setting(self.db,'structure_timezone','America/Montreal')
            self.db.execute("UPDATE fablab_services SET service_date='2026-10-21',end_time='12:00'")
        self.assertEqual(quality.service_statistics(self.db,'total',self.now)['animations'],0)
        self.assertEqual(quality.service_statistics(self.db,'total',self.now.replace(hour=16))['animations'],1)

    def test_latest_slot_end_controls_animation_completion(self):
        from animation_slots import save_slots
        with self.db:
            self.db.execute("UPDATE fablab_services SET service_date='2026-10-21'")
            self.db.execute("UPDATE animation_reservation_config SET booking_mode='slots',slot_duration_minutes=60,slot_capacity=2")
            save_slots(self.db,self.db.execute('SELECT * FROM fablab_services').fetchone(),self.db.execute('SELECT * FROM animation_reservation_config').fetchone(),'Europe/Paris')
        self.assertEqual(quality.service_statistics(self.db,'total',self.now.replace(hour=9,minute=30))['animations'],0)
        self.assertEqual(quality.service_statistics(self.db,'total',self.now)['animations'],1)

    def test_missing_end_counts_only_after_day_finished(self):
        with self.db:self.db.execute("UPDATE fablab_services SET service_date='2026-10-21',end_time=NULL")
        self.assertEqual(quality.service_statistics(self.db,'total',self.now)['animations'],0)
        self.assertEqual(quality.service_statistics(self.db,'total',self.now.replace(day=22))['animations'],1)

    def test_anonymous_and_named_not_counted_twice_and_waiting_excluded(self):
        booking=self.f.book([self.f.a,self.f.b],staff=True)
        attendance.manual(self.db,self.service,booking['booking_uuids'][0],'present','admin')
        attendance.set_walkins(self.db,self.service,3,'admin')
        self.past(4);value=self.stats()['service_statistics']
        self.assertEqual((value['animation_participants'],value['animation_anonymous'],value['animation_unknown_bookings']),(4,3,1))

    def test_cancelled_bookings_do_not_add_participation(self):
        booking=self.f.book([self.f.a],staff=True)
        with self.db:self.db.execute("UPDATE animation_bookings SET status='cancelled' WHERE external_uuid=?",(booking['booking_uuids'][0],))
        self.past(0);self.assertEqual(self.stats()['service_statistics']['animation_participants'],0)

    def test_no_event_cancellation_flag_in_model_is_not_invented(self):
        columns={r[1] for r in self.db.execute('PRAGMA table_info(fablab_services)')}
        self.assertFalse({'cancelled','held','status'}&columns)

    def add_sessions(self):
        with self.db:
            self.db.execute("INSERT INTO sessions(user_id,check_in,check_out,entry_method) VALUES(?,'2025-08-14T10:00:00+00:00','2025-08-14T12:00:00+00:00','id')",(self.f.a,))
            self.db.execute("INSERT INTO visitors(created_at) VALUES('2025-09-01T10:00:00+00:00')")

    def test_coverage_frequency_and_animations_are_separate(self):
        self.add_sessions();self.past();coverage=self.stats()['coverage']
        self.assertEqual((coverage['first'],coverage['sessions_first'],coverage['visitors_first']),('2025-08-14','2025-08-14','2025-09-01'))
        self.assertTrue(coverage['partial']);self.assertIn('14 août 2025',coverage['message'])

    def test_no_frequency_is_not_start_of_software(self):
        self.past();self.assertIsNone(self.stats()['coverage']['first'])
        self.assertEqual(self.stats()['coverage']['animations_first'],'2025-08-14')

    def test_empty_period_and_empty_total(self):
        value=self.stats(2020);self.assertEqual(value['selected_summary']['total'],0)
        self.assertIsNone(value['coverage']['selected_first']);self.assertFalse(value['coverage']['partial'])

    def test_frequency_year_boundary_in_configured_timezone(self):
        with self.db:
            app.write_setting(self.db,'structure_timezone','America/Montreal')
            self.db.execute("INSERT INTO sessions(user_id,check_in,check_out,entry_method) VALUES(?,'2026-01-01T02:00:00+00:00','2026-01-01T03:00:00+00:00','id')",(self.f.a,))
        self.assertEqual(self.stats(2025)['selected_summary']['sessions'],1)
        self.assertEqual(self.stats(2026)['selected_summary']['sessions'],0)

    def test_manual_status_values_and_automatic_priority_preserved(self):
        booking=self.f.book([self.f.a],staff=True)['booking_uuids'][0]
        for action,value in [('present',1),('absent',0),('unknown',None)]:
            attendance.manual(self.db,self.service,booking,action,'moderator')
            self.assertEqual(self.db.execute('SELECT is_present FROM animation_bookings WHERE external_uuid=?',(booking,)).fetchone()[0],value)
            self.assertEqual(self.db.execute('SELECT origin FROM animation_attendance_origin WHERE booking_uuid=?',(booking,)).fetchone()[0],'manual')

    def test_screen_csv_annual_and_visual_projection_share_figures(self):
        self.past(7);self.add_sessions();value=self.stats();self.f.f.login_admin()
        page=self.client.get('/admin/frequentation/statistiques?year=2025').get_data(as_text=True)
        self.assertIn('Animations passées',page);self.assertNotIn('Animations effectuées',page)
        response=self.client.get('/admin/exports/statistiques.csv?year=2025')
        text=response.get_data(as_text=True);self.assertIn('Participations renseignées aux animations passées;7',text)
        self.assertIn(('Participations renseignées',7),_metric_rows(value))
        self.assertEqual(public_summary(value,{'name':'Atelier fictif'})['services']['animation_participants'],7)
        self.assertEqual(value['selected_summary']['hours'],2.0)

    def test_six_csv_exports_and_one_pdf_remain(self):
        self.f.f.login_admin();page=self.client.get('/admin/frequentation/exports').get_data(as_text=True)
        self.assertEqual(len(re.findall(r'href="[^"]+\.csv',page)),6)
        self.assertEqual(page.count('Bilan visuel de fréquentation (PDF)'),1)

    def test_visual_pdf_requires_existing_admin_export_role(self):
        self.assertIn(self.client.get('/admin/exports/bilan-visuel.pdf').status_code,(302,403))
        with self.client.session_transaction() as session:session['access_role']='moderator';session['admin_authenticated']=True
        response=self.client.get('/admin/exports/bilan-visuel.pdf')
        self.assertEqual(response.status_code,302);self.assertFalse(response.data.startswith(b'%PDF'))

    def test_visual_pdf_module_disabled_refused(self):
        self.f.f.login_admin()
        with self.db:app.write_setting(self.db,'module_frequency','0')
        self.assertEqual(self.client.get('/admin/exports/bilan-visuel.pdf').status_code,404)

    def test_visual_pdf_real_response_private_no_store(self):
        self.past(7);self.add_sessions();self.f.f.login_admin()
        response=self.client.get('/admin/exports/bilan-visuel.pdf?year=2025')
        self.assertEqual(response.status_code,200);self.assertTrue(response.data.startswith(b'%PDF'))
        self.assertIn('no-store',response.headers['Cache-Control'])
        text=pdf_text(response.data)
        for unwanted in ('person0@example.invalid','Fictif0','EXEMPLE','2001','0600000000'):
            self.assertNotIn(unwanted,text)
        self.assertIn('14 août 2025',text);self.assertIn('7 participations renseignées',text)

    def test_internal_projection_omits_individuals_but_keeps_small_demographics(self):
        self.add_sessions();value=public_summary(self.stats(),{'name':'Structure fictive'})
        self.assertTrue(all(d['available'] for d in value['distributions']))
        self.assertTrue(all(d['total']==1 for d in value['distributions']))
        encoded=json.dumps(value);self.assertNotIn('first_name',encoded);self.assertNotIn('selected_sessions',encoded)

    def test_small_cell_does_not_hide_whole_dimension(self):
        value=self.stats();value['selected_summary']['unique_users']=20
        value['gender_rows']=[dict(label='A',count=19),dict(label='B',count=1)]
        self.assertTrue(public_summary(value,{})['distributions'][1]['available'])
        value['gender_rows']=[dict(label='A',count=15),dict(label='B',count=5)]
        self.assertTrue(public_summary(value,{})['distributions'][1]['available'])

    def test_annual_texts_are_never_rewritten_by_statistic_read(self):
        with self.db:self.db.execute("INSERT INTO annual_activity_reports(year,introduction,created_at,updated_at) VALUES(2025,'Texte historique fictif','2025-01-01','2025-01-01')")
        before=[tuple(r) for r in self.db.execute('SELECT * FROM annual_activity_reports')]
        self.past();self.stats();self.assertEqual(before,[tuple(r) for r in self.db.execute('SELECT * FROM annual_activity_reports')])

    def test_existing_annual_pdf_docx_use_same_corrected_numbers(self):
        self.past(7);self.add_sessions();report={'year':2025,'introduction':'Exemple fictif'}
        statistics=self.stats();pdf=generate_activity_report_pdf(report,statistics)
        self.assertIn('Animations passées',pdf_text(pdf));self.assertIn('14 août 2025',pdf_text(pdf))
        from docx import Document
        document=Document(io.BytesIO(generate_activity_report_docx(report,statistics)))
        self.assertTrue(any('14 août 2025' in p.text for p in document.paragraphs))
        self.assertIn('Participations renseignées',' '.join(c.text for t in document.tables for r in t.rows for c in r.cells))

    def test_diagnostic_is_read_only_no_private_values(self):
        before=self.db.total_changes;queries=[];self.db.set_trace_callback(queries.append)
        with mock.patch('socket.create_connection',side_effect=AssertionError('network')):value=health.diagnostic(self.db,self.now)
        self.db.set_trace_callback(None);self.assertEqual(before,self.db.total_changes)
        self.assertTrue(all(q.lstrip().upper().startswith(('SELECT','PRAGMA','WITH')) for q in queries))
        self.assertEqual(value['schema'],22);self.assertIn('Aucun résultat',value['integrity'])
        for key in ('password','token','webhook','secret','payload_json'):self.assertNotIn(key,json.dumps(value))

    def test_optional_wordpress_disabled_hidden(self):
        with self.db:app.write_setting(self.db,'module_public_reservations','0')
        self.assertIsNone(health.diagnostic(self.db)['wordpress'])

    def test_unused_wordpress_hidden_even_with_local_reservations(self):
        self.assertIsNone(health.diagnostic(self.db)['wordpress'])

    def test_invalid_saved_timestamps_are_unavailable_not_errors(self):
        with self.db:app.write_setting(self.db,'automatic_backup_last_success','invalid')
        self.assertIsNone(health.diagnostic(self.db)['backup_success'])

    def test_coverage_complete_span_and_observed_empty_year(self):
        sessions=[dict(local_date=__import__('datetime').date(y,m,d)) for y,m,d in [(2024,1,1),(2025,1,1),(2025,12,31),(2026,1,1)]]
        self.assertFalse(quality.coverage(self.db,2025,sessions,[],self.now)['partial'])
        self.assertEqual(quality.coverage(self.db,2025,sessions,[],self.now)['selected_first'],'2025-01-01')
        self.assertIsNone(quality.coverage(self.db,2023,sessions,[],self.now)['selected_first'])

    def test_zero_hour_chart_never_invents_a_maximum_of_one(self):
        from visual_report import _hour_chart
        drawing=_hour_chart(dict(bars=[dict(start='10:00',average_users=0)]))
        self.assertTrue(any('Maximum : 0' in getattr(item,'text','') for item in drawing.contents))

    def test_old_commands_without_sent_date_are_not_relay_pending(self):
        with self.db:
            self.db.execute("INSERT INTO reservation_outbox(environment,command_type,entity_key,payload_json,created_at) VALUES('production','upsert','1','{}','2020-01-01')")
        value=health.queue_state(self.db,'production',self.now)
        self.assertEqual(value['legacy'],1);self.assertIsNone(value['pending'])

    def test_acknowledged_recent_and_old_receipts_are_distinct(self):
        with self.db:
            for index,ack,age in [(1,1,1000),(2,None,15),(3,None,500)]:
                self.db.execute("INSERT INTO wordpress_action_receipts VALUES('production',?,'fake','reserve','{}',?,0,?)",(str(index),int(self.now.timestamp())-age,ack))
        value=health.queue_state(self.db,'production',self.now)
        self.assertEqual((value['acknowledged'],value['awaiting_ack'],value['retry_ack']),(1,1,1))

    def test_relay_configuration_does_not_prove_success(self):
        with self.db:app.write_setting(self.db,'reservation_wordpress_url','https://example.invalid')
        value=health.diagnostic(self.db,self.now)['wordpress'];self.assertTrue(value['configured'])
        self.assertTrue(all(r['freshness']=='unknown' for r in value['states']))

    def test_health_http_only_admin(self):
        with self.client.session_transaction() as session:session['access_role']='moderator';session['admin_authenticated']=True
        response=self.client.get('/admin/reglages/donnees')
        self.assertEqual(response.status_code,302);self.assertNotIn('État de l’installation',response.get_data(as_text=True))
        with self.client.session_transaction() as session:session.clear()
        self.f.f.login_admin();page=self.client.get('/admin/reglages/donnees').get_data(as_text=True)
        self.assertIn('État de l’installation',page)

    def test_health_http_refresh_does_not_write_database(self):
        self.f.f.login_admin()
        def rows():
            return {r[0]:sorted([tuple(v) for v in self.db.execute('SELECT * FROM '+r[0])],key=repr)
                for r in self.db.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}
        before=rows()
        for _ in range(2):self.assertEqual(self.client.get('/admin/reglages/donnees').status_code,200)
        self.assertEqual(before,rows())

def load_tests(loader,suite,pattern):return loader.loadTestsFromTestCase(ConsolidationTests)
